"""File → pages of normalised text, each with a reference ("page:12", "slide:18") so every chunk
can cite where it came from. CPU-bound: the indexing job runs each step in a thread.

PDF pages with almost no text layer (scans, photos of the board) and image files go through
Tesseract OCR (declared service, local)."""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from pathlib import Path

from app.core.config import get_settings

KINDS = {
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "txt",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
}
MIN_TEXT_CHARS = 25  # below this a PDF page is treated as scanned → OCR
OCR_DPI = 200
TXT_PART_CHARS = 3000
WINDOWS_TESSERACT = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")


@dataclass(frozen=True)
class Page:
    ref: str  # "page:12" | "slide:18" | "section:3" | "part:2" | "image:1"
    text: str
    ocr: bool = False


class ParseError(Exception):
    """User-safe failure; `key` is an i18n message key."""

    def __init__(self, key: str) -> None:
        super().__init__(key)
        self.key = key


def kind_for(filename: str) -> str | None:
    return KINDS.get(Path(filename).suffix.lower())


def normalise(text: str) -> str:
    text = unicodedata.normalize("NFC", text).replace("\u00ad", "")  # soft hyphens
    text = re.sub(r"-\n(?=[a-zà-ÿ])", "", text)  # words hyphenated across lines
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# --- OCR ---------------------------------------------------------------------------------------
def tesseract_cmd() -> str | None:
    configured = get_settings().TESSERACT_CMD
    if configured:
        return configured
    found = shutil.which("tesseract")
    if found:
        return found
    return str(WINDOWS_TESSERACT) if WINDOWS_TESSERACT.exists() else None


def ocr_available() -> bool:
    return tesseract_cmd() is not None


def ocr_image(path: Path) -> str:
    cmd = tesseract_cmd()
    if cmd is None:
        return ""
    proc = subprocess.run(
        [cmd, str(path), "stdout", "-l", get_settings().TESSERACT_LANGS],
        capture_output=True,
        timeout=120,
        check=False,
    )
    return proc.stdout.decode("utf-8", errors="replace") if proc.returncode == 0 else ""


def _ocr_png_bytes(png: bytes) -> str:
    fd, name = tempfile.mkstemp(suffix=".png")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(png)
        return ocr_image(Path(name))
    finally:
        os.unlink(name)


# --- per format: (total units, iterator of thunks producing one Page each) ------------------------
PageSteps = tuple[int, Iterator[Callable[[], Page | None]]]


def _pdf(path: Path) -> PageSteps:
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as e:
        raise ParseError("errors.document_unreadable") from e

    def step(i: int) -> Callable[[], Page | None]:
        def run() -> Page | None:
            page = doc[i]
            text = normalise(page.get_text("text", sort=True))
            if len(text) >= MIN_TEXT_CHARS:
                return Page(f"page:{i + 1}", text)
            ocr = normalise(_ocr_png_bytes(page.get_pixmap(dpi=OCR_DPI).tobytes("png")))
            return Page(f"page:{i + 1}", ocr, ocr=True) if ocr else None

        return run

    return doc.page_count, (step(i) for i in range(doc.page_count))


def _pptx(path: Path) -> PageSteps:
    from pptx import Presentation

    try:
        slides = list(Presentation(str(path)).slides)
    except Exception as e:
        raise ParseError("errors.document_unreadable") from e

    def step(i: int) -> Callable[[], Page | None]:
        def run() -> Page | None:
            slide = slides[i]
            parts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
            text = normalise("\n".join(p for p in parts if p.strip()))
            return Page(f"slide:{i + 1}", text) if text else None

        return run

    return len(slides), (step(i) for i in range(len(slides)))


def _docx(path: Path) -> PageSteps:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as e:
        raise ParseError("errors.document_unreadable") from e
    # Word has no pages: a section starts at each heading.
    sections: list[list[str]] = [[]]
    for para in document.paragraphs:
        if (
            para.style is not None
            and para.style.name.lower().startswith("heading")
            and sections[-1]
        ):
            sections.append([])
        if para.text.strip():
            sections[-1].append(para.text)
    for table in document.tables:
        sections.append([" | ".join(c.text for c in row.cells) for row in table.rows])
    sections = [s for s in sections if s]

    def step(i: int) -> Callable[[], Page | None]:
        return lambda: Page(f"section:{i + 1}", normalise("\n".join(sections[i])))

    return len(sections), (step(i) for i in range(len(sections)))


def _txt(path: Path) -> PageSteps:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    paras = re.split(r"\n\s*\n", text)
    parts: list[str] = [""]
    for p in paras:
        if parts[-1] and len(parts[-1]) + len(p) > TXT_PART_CHARS:
            parts.append("")
        parts[-1] += p + "\n\n"
    parts = [p for p in parts if p.strip()]

    def step(i: int) -> Callable[[], Page | None]:
        return lambda: Page(f"part:{i + 1}", normalise(parts[i]))

    return len(parts), (step(i) for i in range(len(parts)))


def _image(path: Path) -> PageSteps:
    def run() -> Page | None:
        text = normalise(ocr_image(path))
        return Page("image:1", text, ocr=True) if text else None

    return 1, iter([run])


PARSERS: dict[str, Callable[[Path], PageSteps]] = {
    "pdf": _pdf,
    "pptx": _pptx,
    "docx": _docx,
    "txt": _txt,
    "image": _image,
}


def page_steps(path: Path, kind: str) -> PageSteps:
    """(unit count, lazy steps). The caller runs each step in a thread and reports progress."""
    return PARSERS[kind](path)


def parse(path: Path, kind: str) -> list[Page]:
    _, steps = page_steps(path, kind)
    return [p for step in steps if (p := step()) is not None]
