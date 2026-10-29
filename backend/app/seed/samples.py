"""Synthetic course material written for this project (CC0): a 30-page Calculus I PDF (Limits,
Derivatives; one page rasterised as a "scan"), a Data Structures PPTX (stacks, queues), a Word
exercise sheet and a photo-of-the-board PNG. Used by the seed script and by the tests, so the
demo world never contains third-party or personal material."""

from __future__ import annotations

from pathlib import Path

LIMITS_PARAGRAPHS = [
    "The limit of f(x) as x approaches a is the value that f(x) gets close to when x gets close "
    "to a. The function does not need to be defined at a: the limit only looks at nearby points.",
    "Direct substitution is the first thing to try. If f is a polynomial or a rational function "
    "whose denominator is not zero at a, the limit is simply f(a).",
    "When substitution gives 0/0 the form is indeterminate. An indeterminate form 0/0 does not "
    "mean the limit fails to exist; it means we must simplify before substituting again.",
    "Factor and cancel: (x^2 - 4)/(x - 2) gives 0/0 at x = 2, but it equals x + 2 for x != 2, so "
    "the limit as x approaches 2 is 4.",
    "Rationalise: for (sqrt(x + 1) - 1)/x at x = 0 we get 0/0. Multiply by the conjugate "
    "sqrt(x + 1) + 1 to obtain 1/(sqrt(x + 1) + 1), whose limit at 0 is 1/2.",
    "One-sided limits: the left limit uses x < a and the right limit uses x > a. The limit exists "
    "only when both one-sided limits exist and are equal.",
    "A vertical asymptote appears when the denominator tends to 0 but the numerator does not: "
    "1/x grows without bound as x approaches 0 from the right.",
    "Limits at infinity describe end behaviour. For a ratio of polynomials, compare the degrees: "
    "equal degrees give the ratio of leading coefficients.",
    "The squeeze theorem: if g(x) <= f(x) <= h(x) near a and g and h have the same limit L at a, "
    "then f has limit L as well. Example: x^2 sin(1/x) tends to 0.",
    "The special limit sin(x)/x tends to 1 as x approaches 0. It is a 0/0 form that cannot be "
    "factored; it is proved with the squeeze theorem.",
    "Continuity at a means three things: f(a) is defined, the limit at a exists, and the two are "
    "equal. Polynomials are continuous everywhere.",
    "Common mistake: cancelling terms instead of factors. (x^2 - 9)/(x - 3) must be factored as "
    "(x - 3)(x + 3)/(x - 3) before cancelling, giving the limit 6 at x = 3.",
]

DERIVATIVE_PARAGRAPHS = [
    "The derivative f'(a) is the slope of the tangent line to the graph of f at the point "
    "(a, f(a)). It measures the instantaneous rate of change.",
    "Definition: f'(a) is the limit of the difference quotient (f(a + h) - f(a))/h as h tends "
    "to 0. The derivative function f' gives the slope at every point.",
    "Power rule: the derivative of x^n is n x^(n - 1). For example the derivative of x^5 is "
    "5x^4 and the derivative of a constant is 0.",
    "Sum and constant multiple rules let us differentiate term by term: (3x^2 + 2x)' = 6x + 2.",
    "Product rule: (u v)' = u' v + u v'. For x^2 sin(x) the derivative is 2x sin(x) + x^2 cos(x).",
    "Quotient rule: (u / v)' = (u' v - u v') / v^2, valid where v is not zero.",
    "Chain rule: the derivative of f(g(x)) is f'(g(x)) g'(x). For (3x + 1)^4 we get "
    "4(3x + 1)^3 times 3.",
    "Derivatives of trigonometric functions: sin' = cos and cos' = -sin. The tangent line at a "
    "point is y = f(a) + f'(a)(x - a).",
    "A function that is differentiable at a is continuous at a, but |x| is continuous at 0 and "
    "has no derivative there because of the corner.",
    "Critical points are where f'(x) = 0 or f' is undefined. A sign change of f' from + to - "
    "marks a local maximum; from - to + a local minimum.",
    "Rates of change in physics: if s(t) is position then s'(t) is velocity and s''(t) is "
    "acceleration.",
    "Common mistake: forgetting the inner derivative in the chain rule, writing 4(3x + 1)^3 "
    "instead of 12(3x + 1)^3.",
]

STACK_SLIDES = [
    ("Stacks", "A stack is a LIFO structure: last in, first out. Think of a pile of plates."),
    ("Stack operations", "push(x) adds on top, pop() removes the top, peek() reads it. All O(1)."),
    ("Array-based stack", "Keep an array and a top index. push writes at top and increments."),
    ("Linked stack", "Each push creates a node pointing to the previous top; no fixed capacity."),
    ("Uses of stacks", "Undo history, the call stack, checking balanced brackets, DFS."),
    ("Balanced brackets", "Push every opening bracket; on a closing one pop and compare."),
    ("Queues", "A queue is FIFO: first in, first out, like a line at the bakery."),
    ("Queue operations", "enqueue(x) adds at the back, dequeue() removes from the front. O(1)."),
    ("Circular buffer", "An array with head and tail indices that wrap around modulo capacity."),
    ("Linked queue", "Keep pointers to head and tail nodes; enqueue at tail, dequeue at head."),
    ("Uses of queues", "Printer jobs, BFS on graphs, task scheduling, buffering keystrokes."),
    ("Stack vs queue", "Stack: LIFO, one end. Queue: FIFO, two ends. Deque: both ends."),
]

EXERCISES = [
    ("Exercise sheet 3 · Derivatives", ""),
    ("1. Power rule", "Differentiate f(x) = 4x^3 - 2x + 7. Then find the slope at x = 1."),
    ("2. Product rule", "Differentiate g(x) = x^2 cos(x)."),
    ("3. Chain rule", "Differentiate h(x) = (2x - 5)^6. Do not forget the inner derivative."),
    ("4. Tangent line", "Find the tangent line to y = x^3 at x = 2."),
]

BOARD_TEXT = "Integrals\nArea under the curve\nF'(x) = f(x)\nIntegral of x^2 dx = x^3/3 + C"


def _page_text(title: str, paragraphs: list[str], i: int) -> str:
    # Each page: a heading and 6 paragraphs, rotating through the chapter so pages differ.
    picked = [paragraphs[(i * 3 + k) % len(paragraphs)] for k in range(6)]
    return f"{title} - page {i + 1}\n\n" + "\n\n".join(picked)


def make_pdf(path: Path, pages: int = 30, scanned_page: int | None = 4) -> Path:
    """Pages 1..pages/2 are Chapter 1 Limits, the rest Chapter 2 Derivatives."""
    import pymupdf

    doc = pymupdf.open()
    half = pages // 2
    for i in range(pages):
        limits = i < half
        title = "Chapter 1 - Limits" if limits else "Chapter 2 - Derivatives"
        paras = LIMITS_PARAGRAPHS if limits else DERIVATIVE_PARAGRAPHS
        text = _page_text(title, paras, i if limits else i - half)
        page = doc.new_page(width=595, height=842)
        page.insert_textbox(pymupdf.Rect(50, 50, 545, 800), text, fontsize=11, fontname="helv")
        if scanned_page is not None and i == scanned_page:
            # Rasterise this page and replace it by the picture: no text layer, like a scan.
            pix = page.get_pixmap(dpi=150)
            doc.delete_page(i)
            scan = doc.new_page(pno=i, width=595, height=842)
            scan.insert_image(scan.rect, pixmap=pix)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    for title, body in STACK_SLIDES:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def make_docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    for heading, body in EXERCISES:
        document.add_heading(heading, level=1 if not body else 2)
        if body:
            document.add_paragraph(body)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def make_board_png(path: Path) -> Path:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=420, height=220)
    page.insert_textbox(pymupdf.Rect(20, 20, 400, 200), BOARD_TEXT, fontsize=20, fontname="helv")
    path.parent.mkdir(parents=True, exist_ok=True)
    page.get_pixmap(dpi=150).save(str(path))
    return path


def make_all(folder: Path) -> dict[str, Path]:
    return {
        "pdf": make_pdf(folder / "Calculus_I_Limits_Derivatives.pdf"),
        "pptx": make_pptx(folder / "DataStructures_Stacks_Queues.pptx"),
        "docx": make_docx(folder / "TD3_derivatives_exercises.docx"),
        "png": make_board_png(folder / "board_integrals.png"),
    }
